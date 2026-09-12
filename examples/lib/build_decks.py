#!/usr/bin/env python3
"""Build styled PPTX + self-contained HTML slide decks from the demo
presentation markdown.

Single source of truth: each ``examples/<demo>/presentation/presentation.md``
(house style — ``## SLIDE N``, ``**Title:**``, ``**Bullets:**``, ``**Visual:**``,
``**Speaker Notes (m:ss–m:ss):**`` blockquote). This script parses that
structure and emits, next to each source file:

    presentation.pptx   16:9 deck, speaker notes in the notes pane
    presentation.html   self-contained deck (arrow-key nav, 'n' toggles notes)

Run:  python3.13 examples/lib/build_decks.py
"""
from __future__ import annotations

import html
import re
from dataclasses import dataclass, field
from pathlib import Path

from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE

# ── Palette (deliberate dark deck aesthetic) ───────────────────────────────
BG        = RGBColor(0x0B, 0x12, 0x20)
PANEL     = RGBColor(0x0F, 0x17, 0x2A)
CODE_BG   = RGBColor(0x11, 0x18, 0x27)
ACCENT    = RGBColor(0x22, 0xD3, 0xEE)   # cyan — matches the CLI '»'
ACCENT2   = RGBColor(0xC0, 0x84, 0xFC)   # violet — used on the closing slide
TEXT      = RGBColor(0xE2, 0xE8, 0xF0)
MUTED     = RGBColor(0x94, 0xA3, 0xB8)
CODE_TEXT = RGBColor(0xE5, 0xE7, 0xEB)

HEX = dict(bg="#0B1220", panel="#0F172A", codebg="#0d1424", accent="#22D3EE",
           accent2="#C084FC", text="#E2E8F0", muted="#94A3B8", codetext="#E5E7EB",
           border="#1E293B")

REPO = Path(__file__).resolve().parents[2]
DECKS = [
    REPO / "examples/01-quickstart/presentation/presentation.md",
    REPO / "examples/02-resident-session/presentation/presentation.md",
    REPO / "examples/03-sales-report/presentation/presentation.md",
]


# ── Parse ──────────────────────────────────────────────────────────────────
@dataclass
class Slide:
    kind: str = "content"          # "title" | "content" | "close"
    tag: str = ""                  # "SLIDE 3 — THE DEMO"
    title: str = ""
    subtitle: str = ""
    bullets: list[str] = field(default_factory=list)
    ordered: bool = False
    code: str = ""
    visual: str = ""
    timing: str = ""
    notes: str = ""


@dataclass
class Deck:
    title: str = ""
    kicker: str = ""               # "5-Minute Demo Presentation — Demo 1"
    brand: str = ""
    slides: list[Slide] = field(default_factory=list)


def parse(md: str) -> Deck:
    blocks = re.split(r"(?m)^---\s*$", md)
    head, *rest = blocks

    deck = Deck()
    m = re.search(r"(?m)^#\s+(.+)$", head)
    deck.title = m.group(1).strip() if m else "Presentation"
    m = re.search(r"(?m)^\*\*(.+?)\*\*", head)
    deck.kicker = m.group(1).strip() if m else ""
    m = re.search(r"(?m)^\*(.+?)\*\s*$", head)
    deck.brand = m.group(1).strip() if m else "Kiraa AI / SwiftPandas"

    for block in rest:
        if "## SLIDE" not in block:
            continue
        deck.slides.append(parse_slide(block))
    # Mark first as title card, last as close (for accent styling).
    if deck.slides:
        deck.slides[0].kind = "title"
        deck.slides[-1].kind = "close" if deck.slides[-1].kind != "title" else "title"
    return deck


def parse_slide(block: str) -> Slide:
    s = Slide()
    tag = re.search(r"(?m)^##\s+(.+)$", block)
    s.tag = tag.group(1).strip() if tag else ""

    def field_after(label: str) -> str:
        m = re.search(rf"(?m)^\*\*{label}:\*\*\s*(.+?)\s*$", block)
        return m.group(1).strip() if m else ""

    s.title = field_after("Title")
    s.subtitle = field_after("Subtitle")
    s.visual = field_after("Visual")

    tm = re.search(r"\*\*Speaker Notes\s*\(([^)]*)\)", block)
    s.timing = tm.group(1).strip() if tm else ""

    # Bullets: the list that follows the **Bullets:** marker.
    bm = re.search(r"(?m)^\*\*Bullets:\*\*\s*$", block)
    if bm:
        tail = block[bm.end():]
        for line in tail.splitlines():
            if re.match(r"^\s*(?:-|\d+\.)\s+", line):
                if re.match(r"^\s*\d+\.", line):
                    s.ordered = True
                s.bullets.append(re.sub(r"^\s*(?:-|\d+\.)\s+", "", line).strip())
            elif line.strip().startswith("**") or line.strip().startswith("```"):
                break
            elif line.strip() == "" and s.bullets:
                break

    code = re.search(r"```[a-zA-Z]*\n(.*?)```", block, re.S)
    if code:
        s.code = code.group(1).rstrip("\n")

    # Speaker notes: the blockquote after the Speaker Notes marker.
    nm = re.search(r"\*\*Speaker Notes[^\n]*\n((?:>.*\n?)+)", block)
    if nm:
        quote = "\n".join(re.sub(r"^>\s?", "", ln) for ln in nm.group(1).splitlines())
        s.notes = quote.strip().strip('"').strip("“”").strip()
    return s


# ── Inline formatting ──────────────────────────────────────────────────────
def strip_md(text: str) -> str:
    text = re.sub(r"\*\*(.+?)\*\*", r"\1", text)
    text = re.sub(r"`(.+?)`", r"\1", text)
    return text


def html_inline(text: str) -> str:
    out = html.escape(text)
    out = re.sub(r"\*\*(.+?)\*\*", r"<strong>\1</strong>", out)
    out = re.sub(r"`(.+?)`", r"<code>\1</code>", out)
    return out


# ── PPTX ───────────────────────────────────────────────────────────────────
W, H = Inches(13.333), Inches(7.5)


def rect(slide, x, y, w, h, color, line=None):
    shp = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, x, y, w, h)
    shp.fill.solid()
    shp.fill.fore_color.rgb = color
    if line is None:
        shp.line.fill.background()
    else:
        shp.line.color.rgb = line
        shp.line.width = Pt(1)
    shp.shadow.inherit = False
    return shp


def textbox(slide, x, y, w, h, anchor=MSO_ANCHOR.TOP):
    tb = slide.shapes.add_textbox(x, y, w, h)
    tf = tb.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    return tf


def set_run(p, text, size, color, bold=False, mono=False):
    r = p.add_run()
    r.text = text
    r.font.size = Pt(size)
    r.font.color.rgb = color
    r.font.bold = bold
    r.font.name = "Menlo" if mono else "Helvetica Neue"


def build_pptx(deck: Deck, out: Path) -> None:
    prs = Presentation()
    prs.slide_width, prs.slide_height = W, H
    blank = prs.slide_layouts[6]

    for idx, s in enumerate(deck.slides):
        slide = prs.slides.add_slide(blank)
        slide.background.fill.solid()
        slide.background.fill.fore_color.rgb = BG
        accent = ACCENT2 if s.kind == "close" else ACCENT

        if s.kind == "title":
            rect(slide, 0, Inches(3.25), Inches(0.9), Inches(0.09), accent)  # accent tick
            tf = textbox(slide, Inches(0.9), Inches(2.2), Inches(11.5), Inches(2.2))
            p = tf.paragraphs[0]
            set_run(p, s.title, 46, TEXT, bold=True)
            if s.subtitle:
                tf2 = textbox(slide, Inches(0.92), Inches(4.0), Inches(11.2), Inches(1.6))
                set_run(tf2.paragraphs[0], strip_md(s.subtitle), 20, MUTED)
            tf3 = textbox(slide, Inches(0.92), Inches(6.6), Inches(11.2), Inches(0.6))
            set_run(tf3.paragraphs[0], f"{deck.kicker}   ·   {deck.brand}", 12, accent)
        else:
            # header
            rect(slide, Inches(0.9), Inches(1.35), Inches(0.55), Inches(0.07), accent)
            th = textbox(slide, Inches(0.9), Inches(0.55), Inches(11.5), Inches(0.9))
            set_run(th.paragraphs[0], s.title, 30, TEXT, bold=True)
            tk = textbox(slide, Inches(0.9), Inches(0.2), Inches(11.5), Inches(0.4))
            set_run(tk.paragraphs[0], s.tag, 11, accent, bold=True)

            body_top = Inches(1.7)
            body_h = Inches(4.4)
            if s.code:
                # bullets left, code right
                bw = Inches(6.4)
                add_bullets(slide, s, Inches(0.9), body_top, bw, body_h, accent)
                cx, cw = Inches(7.5), Inches(4.9)
                box = rect(slide, cx, body_top, cw, Inches(2.6), CODE_BG, line=RGBColor(0x1E,0x29,0x3B))
                ctf = box.text_frame
                ctf.word_wrap = True
                ctf.margin_left = Inches(0.2); ctf.margin_right = Inches(0.2)
                ctf.margin_top = Inches(0.15); ctf.margin_bottom = Inches(0.15)
                first = True
                for line in s.code.splitlines():
                    p = ctf.paragraphs[0] if first else ctf.add_paragraph()
                    first = False
                    set_run(p, line if line else " ", 11, CODE_TEXT, mono=True)
            else:
                add_bullets(slide, s, Inches(0.9), body_top, Inches(11.4), body_h, accent)

            if s.visual:
                vf = textbox(slide, Inches(0.9), Inches(6.55), Inches(9.5), Inches(0.7))
                set_run(vf.paragraphs[0], "▨ " + strip_md(s.visual), 11, MUTED)
            if s.timing:
                tfm = textbox(slide, Inches(11.0), Inches(6.55), Inches(1.9), Inches(0.5))
                pt = tfm.paragraphs[0]; pt.alignment = PP_ALIGN.RIGHT
                set_run(pt, s.timing, 11, accent)

        # slide number
        sn = textbox(slide, Inches(12.2), Inches(0.15), Inches(1.0), Inches(0.4))
        pn = sn.paragraphs[0]; pn.alignment = PP_ALIGN.RIGHT
        set_run(pn, f"{idx+1:02d}/{len(deck.slides):02d}", 10, MUTED)

        if s.notes:
            note = slide.notes_slide.notes_text_frame
            note.text = (f"[{s.timing}] " if s.timing else "") + s.notes

    prs.save(str(out))


def add_bullets(slide, s: Slide, x, y, w, h, accent) -> None:
    tf = textbox(slide, x, y, w, h)
    for i, b in enumerate(s.bullets):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.space_after = Pt(10)
        marker = f"{i+1}.  " if s.ordered else "▸  "
        set_run(p, marker, 18, accent, bold=True)
        # inline: split bold/code spans into runs
        for text, kind in inline_runs(b):
            set_run(p, text, 18, TEXT, bold=(kind == "b"), mono=(kind == "c"))


def inline_runs(text: str):
    """Yield (text, kind) where kind is '' | 'b' (bold) | 'c' (code)."""
    tokens = re.split(r"(\*\*.+?\*\*|`.+?`)", text)
    for tok in tokens:
        if not tok:
            continue
        if tok.startswith("**") and tok.endswith("**"):
            yield tok[2:-2], "b"
        elif tok.startswith("`") and tok.endswith("`"):
            yield tok[1:-1], "c"
        else:
            yield tok, ""


# ── HTML ───────────────────────────────────────────────────────────────────
def build_html(deck: Deck, out: Path) -> None:
    sections = []
    n = len(deck.slides)
    for i, s in enumerate(deck.slides):
        accent_var = "var(--accent2)" if s.kind == "close" else "var(--accent)"
        body = []
        if s.kind == "title":
            body.append(f'<div class="tick" style="background:{accent_var}"></div>')
            body.append(f'<h1>{html_inline(s.title)}</h1>')
            if s.subtitle:
                body.append(f'<p class="subtitle">{html_inline(s.subtitle)}</p>')
            body.append(f'<p class="brand" style="color:{accent_var}">{html.escape(deck.kicker)} &nbsp;·&nbsp; {html.escape(deck.brand)}</p>')
        else:
            body.append(f'<div class="tag" style="color:{accent_var}">{html.escape(s.tag)}</div>')
            body.append(f'<h2 style="--a:{accent_var}">{html_inline(s.title)}</h2>')
            cols_open = '<div class="cols">' if s.code else '<div class="single">'
            body.append(cols_open)
            if s.bullets:
                tag = "ol" if s.ordered else "ul"
                items = "".join(f"<li>{html_inline(b)}</li>" for b in s.bullets)
                body.append(f'<{tag} style="--a:{accent_var}">{items}</{tag}>')
            if s.code:
                body.append(f'<pre class="code"><code>{html.escape(s.code)}</code></pre>')
            body.append("</div>")
            foot = []
            if s.visual:
                foot.append(f'<span class="visual">▨ {html_inline(s.visual)}</span>')
            if s.timing:
                foot.append(f'<span class="timing" style="color:{accent_var}">{html.escape(s.timing)}</span>')
            if foot:
                body.append(f'<div class="foot">{"".join(foot)}</div>')
        notes = f'<div class="notes-body">{html_inline(s.notes)}</div>' if s.notes else '<div class="notes-body muted">—</div>'
        sections.append(
            f'<section class="slide {s.kind}" data-index="{i}">'
            f'<div class="content">{"".join(body)}</div>'
            f'<div class="page">{i+1:02d} / {n:02d}</div>'
            f'<div class="notes" hidden><div class="notes-h">Speaker notes'
            f'{f" · {html.escape(s.timing)}" if s.timing else ""}</div>{notes}</div>'
            f'</section>'
        )

    doc = HTML_TEMPLATE.format(
        title=html.escape(deck.title),
        deck_title=html.escape(deck.title),
        brand=html.escape(deck.brand),
        slides="\n".join(sections),
        **HEX,
    )
    out.write_text(doc, encoding="utf-8")


HTML_TEMPLATE = """<!doctype html>
<html lang="en">
<head>
<meta charset="utf-8">
<meta name="viewport" content="width=device-width, initial-scale=1">
<title>{title}</title>
<style>
  :root {{
    --bg:{bg}; --panel:{panel}; --codebg:{codebg}; --accent:{accent};
    --accent2:{accent2}; --text:{text}; --muted:{muted}; --codetext:{codetext};
    --border:{border};
  }}
  * {{ box-sizing:border-box; }}
  html,body {{ margin:0; height:100%; background:#05080f; color:var(--text);
    font-family:-apple-system,"Helvetica Neue",Arial,sans-serif; }}
  .deck {{ position:fixed; inset:0; display:flex; align-items:center; justify-content:center; }}
  .slide {{
    position:absolute; width:min(94vw,1280px); aspect-ratio:16/9; max-height:94vh;
    background:radial-gradient(120% 120% at 12% 0%, #101c30 0%, var(--bg) 55%);
    border:1px solid var(--border); border-radius:18px; padding:5.2% 6%;
    opacity:0; transform:translateY(14px) scale(.99); pointer-events:none;
    transition:opacity .28s ease, transform .28s ease; overflow:hidden;
    box-shadow:0 30px 90px rgba(0,0,0,.55);
  }}
  .slide.active {{ opacity:1; transform:none; pointer-events:auto; }}
  .content {{ height:100%; display:flex; flex-direction:column; }}
  h1 {{ font-size:clamp(28px,4.6vw,60px); line-height:1.05; margin:.1em 0 .2em; letter-spacing:-.02em; }}
  h2 {{ font-size:clamp(22px,3.2vw,40px); margin:.1em 0 .5em; letter-spacing:-.01em;
        padding-bottom:.28em; position:relative; }}
  h2::after {{ content:""; position:absolute; left:0; bottom:0; width:56px; height:5px;
               border-radius:3px; background:var(--a,var(--accent)); }}
  .tick {{ width:64px; height:7px; border-radius:4px; margin-bottom:26px; }}
  .subtitle {{ font-size:clamp(16px,1.9vw,24px); color:var(--muted); max-width:60ch; line-height:1.4; }}
  .brand {{ margin-top:auto; font-size:13px; letter-spacing:.04em; text-transform:uppercase; }}
  .tag {{ font-size:12px; font-weight:700; letter-spacing:.12em; text-transform:uppercase; margin-bottom:.2em; }}
  .cols {{ display:grid; grid-template-columns:1.05fr .95fr; gap:32px; align-items:start; flex:1; min-height:0; }}
  .single {{ flex:1; min-height:0; }}
  ul,ol {{ margin:.2em 0; padding-left:0; list-style:none; }}
  li {{ position:relative; padding-left:1.6em; margin:.52em 0; font-size:clamp(15px,1.7vw,21px); line-height:1.35; }}
  ul li::before {{ content:"▸"; position:absolute; left:0; color:var(--a,var(--accent)); font-weight:700; }}
  ol {{ counter-reset:li; }}
  ol li {{ counter-increment:li; }}
  ol li::before {{ content:counter(li,decimal-leading-zero); position:absolute; left:0;
    color:var(--a,var(--accent)); font-weight:700; font-variant-numeric:tabular-nums; }}
  code {{ font-family:"SF Mono",Menlo,monospace; background:rgba(34,211,238,.1);
    color:var(--codetext); padding:.05em .35em; border-radius:5px; font-size:.9em; }}
  pre.code {{ background:var(--codebg); border:1px solid var(--border); border-radius:12px;
    padding:16px 18px; overflow:auto; margin:0; font-size:clamp(11px,1.15vw,14px); }}
  pre.code code {{ background:none; color:var(--codetext); padding:0; white-space:pre; }}
  .foot {{ margin-top:auto; padding-top:14px; display:flex; justify-content:space-between;
    align-items:flex-end; gap:16px; color:var(--muted); font-size:13px; }}
  .visual {{ max-width:74%; line-height:1.3; }}
  .timing {{ font-variant-numeric:tabular-nums; font-weight:600; white-space:nowrap; }}
  .page {{ position:absolute; top:20px; right:26px; color:var(--muted); font-size:12px;
    font-variant-numeric:tabular-nums; letter-spacing:.06em; }}
  .notes {{ position:absolute; left:0; right:0; bottom:0; background:rgba(5,8,15,.96);
    border-top:1px solid var(--border); padding:16px 22px; max-height:42%; overflow:auto; }}
  .notes-h {{ font-size:11px; letter-spacing:.14em; text-transform:uppercase; color:var(--accent); margin-bottom:6px; }}
  .notes-body {{ font-size:15px; line-height:1.45; color:#cbd5e1; }}
  .muted {{ color:var(--muted); }}
  .hud {{ position:fixed; bottom:16px; left:50%; transform:translateX(-50%); display:flex; gap:10px;
    align-items:center; background:rgba(15,23,42,.8); border:1px solid var(--border);
    border-radius:999px; padding:7px 14px; font-size:12px; color:var(--muted); backdrop-filter:blur(8px); z-index:10; }}
  .hud button {{ background:none; border:0; color:var(--text); cursor:pointer; font-size:14px; padding:2px 6px; border-radius:6px; }}
  .hud button:hover {{ background:rgba(34,211,238,.15); }}
  .hud .sep {{ opacity:.4; }}
  .bar {{ position:fixed; top:0; left:0; height:3px; background:var(--accent); z-index:11; transition:width .28s ease; }}
  @media (max-width:700px) {{ .cols {{ grid-template-columns:1fr; gap:14px; }} .slide {{ padding:6% 6.5%; }} }}
</style>
</head>
<body>
  <div class="bar" id="bar"></div>
  <div class="deck" id="deck">
    {slides}
  </div>
  <div class="hud">
    <button id="prev" title="Previous (←)">◀</button>
    <span id="counter">1 / 1</span>
    <button id="next" title="Next (→)">▶</button>
    <span class="sep">·</span>
    <button id="notesBtn" title="Toggle speaker notes (N)">notes</button>
  </div>
<script>
  const slides = Array.from(document.querySelectorAll('.slide'));
  let i = 0, notesOn = false;
  const bar = document.getElementById('bar');
  const counter = document.getElementById('counter');
  function render() {{
    slides.forEach((s, k) => {{
      s.classList.toggle('active', k === i);
      s.querySelector('.notes').hidden = !notesOn;
    }});
    counter.textContent = (i + 1) + ' / ' + slides.length;
    bar.style.width = ((i + 1) / slides.length * 100) + '%';
    if (location.hash !== '#' + (i + 1)) history.replaceState(null, '', '#' + (i + 1));
  }}
  function go(n) {{ i = Math.max(0, Math.min(slides.length - 1, n)); render(); }}
  document.getElementById('next').onclick = () => go(i + 1);
  document.getElementById('prev').onclick = () => go(i - 1);
  document.getElementById('notesBtn').onclick = () => {{ notesOn = !notesOn; render(); }};
  document.addEventListener('keydown', e => {{
    if (['ArrowRight',' ','PageDown'].includes(e.key)) {{ e.preventDefault(); go(i + 1); }}
    else if (['ArrowLeft','PageUp'].includes(e.key)) {{ e.preventDefault(); go(i - 1); }}
    else if (e.key === 'Home') go(0);
    else if (e.key === 'End') go(slides.length - 1);
    else if (e.key.toLowerCase() === 'n') {{ notesOn = !notesOn; render(); }}
    else if (e.key.toLowerCase() === 'f') {{
      if (!document.fullscreenElement) document.documentElement.requestFullscreen(); else document.exitFullscreen();
    }}
  }});
  const start = parseInt(location.hash.replace('#','')); if (start) go(start - 1); else render();
</script>
</body>
</html>
"""


def main() -> None:
    for src in DECKS:
        deck = parse(src.read_text(encoding="utf-8"))
        pptx_out = src.with_name("presentation.pptx")
        html_out = src.with_name("presentation.html")
        build_pptx(deck, pptx_out)
        build_html(deck, html_out)
        print(f"{src.parent.parent.name}: {len(deck.slides)} slides -> "
              f"{pptx_out.name}, {html_out.name}")


if __name__ == "__main__":
    main()
